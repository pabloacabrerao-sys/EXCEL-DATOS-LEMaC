#!/usr/bin/env python3
"""Consolida los aportes recibidos en datos/mezclas.json.

Uso:
    python3 scripts/consolidar.py aportes/*.json            # agrega aportes al consolidado
    python3 scripts/consolidar.py aportes/ --xlsx           # ademas escribe datos/consolidado.xlsx
    python3 scripts/consolidar.py --revisar                 # solo informa el estado del consolidado

Los aportes son los .json que la planilla web descarga; tambien acepta .xlsx con el
formato de la planilla original (requiere openpyxl).
"""

import argparse
import datetime
import glob
import json
import os
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ESQUEMA = os.path.join(RAIZ, "datos", "esquema.json")
CONSOLIDADO = os.path.join(RAIZ, "datos", "mezclas.json")


def leer_json(ruta):
    with open(ruta, encoding="utf-8") as fh:
        return json.load(fh)


def cargar_esquema():
    esq = leer_json(ESQUEMA)
    por_id = {c["id"]: c for c in esq["columnas"]}
    por_label = {c["label"]: c for c in esq["columnas"]}
    por_col = {c["col"]: c for c in esq["columnas"]}
    return esq, por_id, por_label, por_col


def clave(mezcla):
    """Identifica una mezcla para no duplicarla entre aportes."""
    d = mezcla.get("datos", {})
    return (
        str(d.get("aportado_por", "")).strip().lower(),
        str(d.get("obra_informe", "")).strip().lower(),
        str(mezcla.get("nombre", "")).strip().lower(),
        str(mezcla.get("tipo", "")).strip(),
    )


def normalizar(mezcla, por_id, por_label, por_col, aportante=""):
    datos = {}
    crudos = mezcla.get("datos") or mezcla.get("d") or {}
    for k, v in crudos.items():
        if v is None or str(v).strip() == "":
            continue
        defi = por_id.get(k) or por_label.get(k) or por_col.get(k)
        if defi:
            datos[defi["id"]] = str(v).strip()
    if aportante and not datos.get("aportado_por"):
        datos["aportado_por"] = aportante
    return {
        "nombre": (mezcla.get("nombre") or "").strip() or "Mezcla",
        "tipo": (mezcla.get("tipo") or "").strip(),
        "datos": datos,
    }


def desde_xlsx(ruta, esq, por_label):
    try:
        import openpyxl
    except ImportError:
        sys.exit("Para leer .xlsx hace falta openpyxl: pip install openpyxl")
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb["Mezclas"] if "Mezclas" in wb.sheetnames else wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    i_enc = next((i for i, f in enumerate(filas)
                  if any(str(v).strip() == "Tipo de mezcla" for v in f if v is not None)), None)
    if i_enc is None:
        sys.exit("%s: no se encontro la fila de encabezados" % ruta)
    enc = [str(v).strip() if v is not None else "" for v in filas[i_enc]]
    i_tipo = enc.index("Tipo de mezcla")
    i_nom = enc.index("Mezcla") if "Mezcla" in enc else None
    mezclas = []
    for f in filas[i_enc + 1:]:
        datos = {}
        for k, v in enumerate(f):
            if k >= len(enc) or v is None or str(v).strip() == "":
                continue
            defi = por_label.get(enc[k])
            if defi:
                datos[defi["id"]] = str(v).strip().replace(".", ",")
        tipo = str(f[i_tipo]).strip() if i_tipo < len(f) and f[i_tipo] else ""
        if not tipo and not datos:
            continue
        nombre = str(f[i_nom]).strip() if i_nom is not None and i_nom < len(f) and f[i_nom] else "Mezcla"
        mezclas.append({"nombre": nombre, "tipo": tipo, "datos": datos})
    return {"aportante": "", "mezclas": mezclas}


def faltantes(mezcla, esq, por_id):
    tipo = next((t for t in esq["tipos"] if t["nombre"] == mezcla["tipo"]), None)
    if tipo is None:
        return None
    faltan = []
    for col in esq["columnas"]:
        if tipo["requisitos"].get(col["col"]) == "obligatorio" and not mezcla["datos"].get(col["id"]):
            faltan.append(col["label"])
    return faltan


def expandir(rutas):
    salida = []
    for r in rutas:
        if os.path.isdir(r):
            salida += sorted(glob.glob(os.path.join(r, "*.json")) + glob.glob(os.path.join(r, "*.xlsx")))
        else:
            salida += sorted(glob.glob(r)) or [r]
    return salida


def escribir_xlsx(esq, mezclas):
    try:
        import openpyxl
    except ImportError:
        sys.exit("Para escribir .xlsx hace falta openpyxl: pip install openpyxl")
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mezclas"
    ws.cell(row=1, column=1, value="Elegir el tipo en la columna B")
    for g in esq["grupos"]:
        idx = 4 + next(i for i, c in enumerate(esq["columnas"]) if c["col"] == g["desde"])
        ws.cell(row=1, column=idx, value=g["titulo"])
    for k, h in enumerate(["Mezcla", "Tipo de mezcla", "Estado"], start=1):
        ws.cell(row=2, column=k, value=h)
    for k, c in enumerate(esq["columnas"], start=4):
        ws.cell(row=2, column=k, value=c["label"])
        ws.column_dimensions[get_column_letter(k)].width = min(28, max(10, len(c["label"]) + 2))
    por_id = {c["id"]: c for c in esq["columnas"]}
    for r, m in enumerate(mezclas, start=3):
        faltan = faltantes(m, esq, por_id)
        ws.cell(row=r, column=1, value=m["nombre"])
        ws.cell(row=r, column=2, value=m["tipo"])
        ws.cell(row=r, column=3, value="Tipo no valido" if faltan is None
                else ("Completa" if not faltan else "Faltan %d" % len(faltan)))
        for k, c in enumerate(esq["columnas"], start=4):
            v = m["datos"].get(c["id"])
            if v is None:
                continue
            if c["tipo"] == "numero":
                try:
                    v = float(str(v).replace(",", "."))
                except ValueError:
                    pass
            ws.cell(row=r, column=k, value=v)
    ws.freeze_panes = "D3"
    rq = wb.create_sheet("Requisitos")
    rq.cell(row=1, column=1, value="No modificar")
    for k, h in enumerate(["Mezcla", "Tipo de mezcla", "Estado"], start=1):
        rq.cell(row=2, column=k, value=h)
    for k, c in enumerate(esq["columnas"], start=4):
        rq.cell(row=2, column=k, value=c["label"])
    for r, t in enumerate(esq["tipos"], start=3):
        rq.cell(row=r, column=1, value=t["clave"])
        rq.cell(row=r, column=2, value=t["nombre"])
        for k, c in enumerate(esq["columnas"], start=4):
            req = t["requisitos"].get(c["col"], "no_aplica")
            rq.cell(row=r, column=k,
                    value={"obligatorio": "Obligatorio", "recomendado": "Recomendado"}.get(req, "-"))
    salida = os.path.join(RAIZ, "datos", "consolidado.xlsx")
    wb.save(salida)
    return salida


def main():
    ap = argparse.ArgumentParser(description="Consolida aportes de mezclas asfalticas.")
    ap.add_argument("aportes", nargs="*", help="archivos .json/.xlsx o carpetas con aportes")
    ap.add_argument("--xlsx", action="store_true", help="escribir tambien datos/consolidado.xlsx")
    ap.add_argument("--revisar", action="store_true", help="solo informar el estado del consolidado")
    args = ap.parse_args()

    esq, por_id, por_label, por_col = cargar_esquema()
    cons = leer_json(CONSOLIDADO) if os.path.exists(CONSOLIDADO) else {"version": 1, "mezclas": []}
    existentes = {clave(m) for m in cons["mezclas"]}

    agregadas = repetidas = 0
    for ruta in expandir(args.aportes):
        if not os.path.exists(ruta):
            print("  ! no existe: %s" % ruta)
            continue
        aporte = desde_xlsx(ruta, esq, por_label) if ruta.lower().endswith((".xlsx", ".xlsm")) else leer_json(ruta)
        quien = aporte.get("aportante", "")
        nuevas = [normalizar(m, por_id, por_label, por_col, quien) for m in aporte.get("mezclas", [])]
        n_ag = 0
        for m in nuevas:
            if clave(m) in existentes:
                repetidas += 1
                continue
            existentes.add(clave(m))
            cons["mezclas"].append(m)
            n_ag += 1
        agregadas += n_ag
        print("  + %s: %d mezcla(s), %d agregada(s)" % (os.path.basename(ruta), len(nuevas), n_ag))

    if args.aportes:
        cons["version"] = esq.get("version", 1)
        cons["actualizado"] = datetime.date.today().isoformat()
        with open(CONSOLIDADO, "w", encoding="utf-8") as fh:
            json.dump(cons, fh, ensure_ascii=False, indent=1)
            fh.write("\n")
        print("Consolidado: %d mezcla(s) en total (+%d nuevas, %d repetidas omitidas)."
              % (len(cons["mezclas"]), agregadas, repetidas))

    completas = incompletas = invalidas = 0
    for m in cons["mezclas"]:
        faltan = faltantes(m, esq, por_id)
        if faltan is None:
            invalidas += 1
            print("  ? %s: tipo no valido (%r)" % (m["nombre"], m["tipo"]))
        elif faltan:
            incompletas += 1
            print("  - %s (%s): faltan %d obligatorios: %s%s"
                  % (m["nombre"], m["tipo"], len(faltan), ", ".join(faltan[:5]),
                     ", ..." if len(faltan) > 5 else ""))
        else:
            completas += 1
    print("Estado: %d completas, %d incompletas, %d con tipo no valido." % (completas, incompletas, invalidas))

    if args.xlsx:
        print("Escrito: %s" % escribir_xlsx(esq, cons["mezclas"]))


if __name__ == "__main__":
    main()
